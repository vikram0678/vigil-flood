"""
Stand-Alone Last Slide (Slide 6: RESEARCH AND REFERENCES) Generator for VIGIL-FLOOD
Engineered for SIH 2026 Grand Finale:
- Exactly 1 slide in PPTX format.
- Strictly clean Sans-Serif typography (Segoe UI / modern style, NO Times New Roman, NO underlines).
- High-authority Indian Government Data Portals (.gov.in):
  * ISRO Bhuvan / NRSC
  * NDMA CAP-SACHET
  * Data.gov.in (OGD India)
  * CWC / India-WRIS
  * IMD Mausam
  * GSI Bhukosh
- Working Prototype & Open-Source Links:
  * GitHub: https://github.com/vikram0678/vigil-flood
  * Live Project: https://vigil-flood.onrender.com/
- Scientific Formulations (Kinematic Wave, Infinite Slope Fs, TreeSHAP, GLO-30 DEM)
"""

import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

TEMPLATE_PATH = r"D:\flash_flood\SIH2026-IDEA-Presentation-Format (3).pptx"
OUTPUT_1SLIDE_PATH = r"D:\flash_flood\VIGIL_FLOOD_RESEARCH_REFERENCES_SLIDE.pptx"
FINAL_MASTER_DECK = r"D:\flash_flood\VIGIL_FLOOD_SIH_FINAL_PRESENTATION.pptx"

# -------------------------------------------------------------
# EXECUTIVE COLOR PALETTE
# -------------------------------------------------------------
C_BG_PAGE      = RGBColor(248, 250, 252)  # #F8FAFC
C_NAVY_DARK    = RGBColor(15, 23, 42)      # #0F172A - Deep Slate/Navy
C_SLATE_BODY   = RGBColor(51, 65, 85)      # #334155 - Slate Gray
C_MUTED        = RGBColor(100, 116, 139)   # #64748B - Muted Slate
C_CYAN_ACCENT  = RGBColor(14, 165, 233)    # #0EA5E9 - Sky Blue
C_CYAN_LIGHT   = RGBColor(240, 249, 255)  # #F0F9FF - Cyan Tint
C_EMERALD      = RGBColor(16, 185, 129)    # #10B981 - Emerald Green
C_EMERALD_LT   = RGBColor(240, 253, 244)  # #F0FDF4 - Emerald Tint
C_BORDER_GRAY  = RGBColor(203, 213, 225)  # #CBD5E1 - Thin Card Outline
C_WHITE        = RGBColor(255, 255, 255)  # Pure White

FONT_FAMILY = "Segoe UI"

def set_font(run, size_pt=11, bold=False, color=C_SLATE_BODY):
    run.font.name = FONT_FAMILY
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = color

def create_card(slide, left, top, width, height, fill_color=C_WHITE, border_color=C_BORDER_GRAY, border_width=1.0):
    """Creates a modern rounded rectangle container card with thin outline."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = fill_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(border_width)
    else:
        card.line.fill.background()
    return card

def build_research_slide(slide):
    # Clean template placeholder text shapes
    for shape in list(slide.shapes):
        if shape.has_text_frame:
            t = shape.text_frame.text.strip()
            if "RESEARCH" in t:
                tf = shape.text_frame
                tf.clear()
                p = tf.paragraphs[0]
                p.text = "RESEARCH AND REFERENCES"
                set_font(p.runs[0], size_pt=20, bold=True, color=C_NAVY_DARK)
            elif any(k in t for k in ["Details / Links", "Your Team Name"]):
                if "Your Team Name" not in t:
                    sp_elem = shape._element
                    sp_elem.getparent().remove(sp_elem)

    # Subtitle Header Banner
    sub_badge = create_card(slide, Inches(0.6), Inches(0.85), Inches(8.5), Inches(0.32), fill_color=C_CYAN_LIGHT, border_color=C_CYAN_ACCENT, border_width=1.0)
    tf_sb = sub_badge.text_frame
    p_sb = tf_sb.paragraphs[0]
    p_sb.alignment = PP_ALIGN.LEFT
    p_sb.margin_left = Inches(0.12)
    p_sb.margin_top = Inches(0.04)
    r_sb = p_sb.add_run()
    r_sb.text = "Official Indian Government Portals (.gov.in)  |  Peer-Reviewed Formulations  |  Live Prototype & GitHub"
    set_font(r_sb, size_pt=9.5, bold=True, color=C_CYAN_ACCENT)

    # =========================================================================
    # LEFT CARD (~55% WIDTH): OFFICIAL INDIAN GOVERNMENT DATA SOURCES (.gov.in)
    # =========================================================================
    c_left = create_card(slide, Inches(0.6), Inches(1.25), Inches(6.8), Inches(4.7), fill_color=C_WHITE, border_color=C_BORDER_GRAY, border_width=1.2)
    tf_l = c_left.text_frame
    tf_l.word_wrap = True
    tf_l.margin_left = Inches(0.2)
    tf_l.margin_right = Inches(0.2)
    tf_l.margin_top = Inches(0.15)

    # Title Tag
    tag_l = create_card(slide, Inches(0.75), Inches(1.32), Inches(4.2), Inches(0.3), fill_color=C_NAVY_DARK, border_color=None)
    tf_tl = tag_l.text_frame
    p_tl = tf_tl.paragraphs[0]
    p_tl.alignment = PP_ALIGN.CENTER
    p_tl.margin_top = Inches(0.04)
    r_tl = p_tl.add_run()
    r_tl.text = "🏛️ OFFICIAL INDIAN GOVT DATA SOURCES (.GOV.IN)"
    set_font(r_tl, size_pt=9.5, bold=True, color=C_WHITE)

    gov_sources = [
        ("ISRO Bhuvan (National Landslide Atlas):", "bhuvan-app1.nrsc.gov.in/landslide/", "Himalayan basin landslide susceptibility zonation & historical scar maps"),
        ("NDMA CAP-SACHET (Disaster Warning):", "sachet.ndma.gov.in/", "ITU-T X.1303 Common Alerting Protocol citizen emergency broadcast standard"),
        ("Data.gov.in (Open Government Data OGD):", "www.data.gov.in/", "National open spatial datasets & district-level meteorological baselines"),
        ("CWC / India-WRIS (River Water Levels):", "indiawris.gov.in/wris/", "Central Water Commission telemetry flood stage thresholds & hydrographs"),
        ("IMD Doppler Radar & Mausam:", "mausam.imd.gov.in/", "India Meteorological Department cloudburst Doppler radar rain rates"),
        ("GSI Bhukosh (Geological Inventories):", "bhukosh.gsi.gov.in/", "Geological Survey of India tectonic fault lines & slope lithology data")
    ]

    p_spacer = tf_l.paragraphs[0]
    p_spacer.space_after = Pt(22)

    for title, url, desc in gov_sources:
        p = tf_l.add_paragraph()
        p.space_after = Pt(4)

        r1 = p.add_run()
        r1.text = "• " + title + " "
        set_font(r1, size_pt=9.5, bold=True, color=C_NAVY_DARK)

        r2 = p.add_run()
        r2.text = url + "\n"
        set_font(r2, size_pt=9, bold=True, color=C_CYAN_ACCENT)

        r3 = p.add_run()
        r3.text = "   " + desc
        set_font(r3, size_pt=8, bold=False, color=C_MUTED)

    # =========================================================================
    # RIGHT TOP CARD (~42% WIDTH): WORKING PROTOTYPE & OPEN-SOURCE DEMO
    # =========================================================================
    c_proj = create_card(slide, Inches(7.55), Inches(1.25), Inches(5.15), Inches(2.15), fill_color=C_CYAN_LIGHT, border_color=C_CYAN_ACCENT, border_width=1.5)
    tf_p = c_proj.text_frame
    tf_p.word_wrap = True
    tf_p.margin_left = Inches(0.2)
    tf_p.margin_right = Inches(0.2)
    tf_p.margin_top = Inches(0.12)

    tag_p = create_card(slide, Inches(7.7), Inches(1.32), Inches(3.6), Inches(0.28), fill_color=C_CYAN_ACCENT, border_color=None)
    tf_tp = tag_p.text_frame
    p_tp = tf_tp.paragraphs[0]
    p_tp.alignment = PP_ALIGN.CENTER
    p_tp.margin_top = Inches(0.03)
    r_tp = p_tp.add_run()
    r_tp.text = "🚀 PROJECT LINKS & WORKING PROTOTYPE"
    set_font(r_tp, size_pt=9, bold=True, color=C_WHITE)

    p_psp = tf_p.paragraphs[0]
    p_psp.space_after = Pt(20)

    proj_links = [
        ("Live Web Application Demo:", "https://vigil-flood.onrender.com/", "Fully deployed 2D/3D tactical digital twin with live simulation sandbox"),
        ("GitHub Open-Source Code:", "https://github.com/vikram0678/vigil-flood", "Complete FastAPI backend, ML models, WebSockets, and React frontend"),
        ("Problem Statement Alignment:", "SIH 2026 | ID: 26192", "Flash Flood Prediction System for Hilly Regions using Multi-Source Data")
    ]

    for title, link, desc in proj_links:
        p = tf_p.add_paragraph()
        p.space_after = Pt(3)

        r1 = p.add_run()
        r1.text = "✔ " + title + " "
        set_font(r1, size_pt=9, bold=True, color=C_NAVY_DARK)

        r2 = p.add_run()
        r2.text = link + "\n"
        set_font(r2, size_pt=8.5, bold=True, color=C_CYAN_ACCENT)

        r3 = p.add_run()
        r3.text = "   " + desc
        set_font(r3, size_pt=7.5, bold=False, color=C_MUTED)

    # =========================================================================
    # RIGHT BOTTOM CARD (~42% WIDTH): SCIENTIFIC FORMULATIONS & ML REFERENCES
    # =========================================================================
    c_sci = create_card(slide, Inches(7.55), Inches(3.55), Inches(5.15), Inches(2.4), fill_color=C_WHITE, border_color=C_BORDER_GRAY, border_width=1.2)
    tf_s = c_sci.text_frame
    tf_s.word_wrap = True
    tf_s.margin_left = Inches(0.2)
    tf_s.margin_right = Inches(0.2)
    tf_s.margin_top = Inches(0.12)

    tag_s = create_card(slide, Inches(7.7), Inches(3.62), Inches(3.6), Inches(0.28), fill_color=C_NAVY_DARK, border_color=None)
    tf_ts = tag_s.text_frame
    p_ts = tf_ts.paragraphs[0]
    p_ts.alignment = PP_ALIGN.CENTER
    p_ts.margin_top = Inches(0.03)
    r_ts = p_ts.add_run()
    r_ts.text = "🔬 SCIENTIFIC FORMULATIONS & AI PAPERS"
    set_font(r_ts, size_pt=9, bold=True, color=C_WHITE)

    p_ssp = tf_s.paragraphs[0]
    p_ssp.space_after = Pt(20)

    sci_items = [
        ("Hydrodynamic Kinematic Wave Routing:", "Lighthill & Whitham surge propagation down bedrock channels: ∂A/∂t + ∂Q/∂x = q"),
        ("Geotechnical Infinite Slope Factor of Safety:", "Limit equilibrium slope stability model under transient pore-water pressure (Fs < 1.0)"),
        ("TreeSHAP Explainable AI (Lundberg et al.):", "Nature Machine Intelligence (2020) game-theoretic Shapley feature contribution"),
        ("Copernicus 30m Digital Elevation Model:", "Copernicus GLO-30 & OpenTopography watershed flow accumulation (D8)")
    ]

    for title, desc in sci_items:
        p = tf_s.add_paragraph()
        p.space_after = Pt(3)

        r1 = p.add_run()
        r1.text = "• " + title + " "
        set_font(r1, size_pt=9, bold=True, color=C_NAVY_DARK)

        r2 = p.add_run()
        r2.text = desc
        set_font(r2, size_pt=8, bold=False, color=C_SLATE_BODY)

    # Bottom Verification Ribbon
    ribbon = create_card(slide, Inches(0.6), Inches(6.08), Inches(12.1), Inches(0.4), fill_color=C_NAVY_DARK, border_color=None)
    tf_r = ribbon.text_frame
    p_r = tf_r.paragraphs[0]
    p_r.alignment = PP_ALIGN.CENTER
    r_r = p_r.add_run()
    r_r.text = "🏛️ 100% OPEN-ACCESS STANDARDS:   NDMA CAP-SACHET Guidelines (2022)   |   ISRO Landslide Atlas   |   Data.gov.in"
    set_font(r_r, size_pt=9.5, bold=True, color=C_WHITE)

def main():
    prs = Presentation(TEMPLATE_PATH)
    print(f"[*] Loaded template with {len(prs.slides)} slides.")

    # In template, slide 6 is index 5
    slide6 = prs.slides[5]
    build_research_slide(slide6)

    # Delete all other slides so that ONLY this slide remains
    # Template has 7 slides (0 to 6). Slide 6 is index 5.
    # Delete slide 7 (index 6)
    if len(prs.slides) > 6:
        rId = prs.slides._sldIdLst[6].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[6]

    # Delete slides 0, 1, 2, 3, 4 (which are indices 0, 1, 2, 3, 4)
    # When deleting index 0 repeatedly, we delete 5 slides
    for _ in range(5):
        rId = prs.slides._sldIdLst[0].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[0]

    prs.save(OUTPUT_1SLIDE_PATH)
    print(f"[SUCCESS] Standalone 1-Slide Research Presentation saved to: {OUTPUT_1SLIDE_PATH}")
    print(f"Total Slides: {len(prs.slides)}")

    # Also update Slide 6 in the final master presentation
    if os.path.exists(FINAL_MASTER_DECK):
        prs_master = Presentation(FINAL_MASTER_DECK)
        if len(prs_master.slides) >= 6:
            slide_m6 = prs_master.slides[5]
            for shape in list(slide_m6.shapes):
                if shape.name not in ["Rectangle 9", "Title 1", "Slide Number Placeholder 5", "Footer Placeholder 6", "Oval 8", "Picture 11"]:
                    sp_elem = shape._element
                    sp_elem.getparent().remove(sp_elem)
            build_research_slide(slide_m6)
            prs_master.save(FINAL_MASTER_DECK)
            print(f"[SUCCESS] Synchronized Slide 6 in Master Deck: {FINAL_MASTER_DECK}")

if __name__ == "__main__":
    main()
