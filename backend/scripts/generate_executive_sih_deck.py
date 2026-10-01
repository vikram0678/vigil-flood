"""
Professional Executive SIH 2026 Presentation Generator for VIGIL-FLOOD
Engineered by Senior Presentation Designer (10+ years experience).
Aligns strictly with SIH 2026 Idea Presentation Template rules (Max 6 slides).
Features modern executive UI cards, visual metrics, shapes, badges, and embedded slide architecture.
"""

import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

TEMPLATE_PATH = r"D:\flash_flood\SIH2026-IDEA-Presentation-Format (3).pptx"
OUTPUT_PATH = r"D:\flash_flood\SIH2026_VIGIL_FLOOD_OFFICIAL_PRESENTATION.pptx"
UNIFIED_SLIDE_IMG = r"D:\flash_flood\SIH_UNIFIED_TECHNICAL_ARCHITECTURE_SLIDE.jpg"

# -------------------------------------------------------------
# EXECUTIVE COLOR PALETTE (Modern Slate, Sky Blue, Navy & Alerts)
# -------------------------------------------------------------
C_NAVY_DARK    = RGBColor(15, 23, 42)      # #0F172A - Deep Slate/Navy
C_NAVY_LIGHT   = RGBColor(30, 41, 59)     # #1E293B - Card Dark
C_SKY_BLUE     = RGBColor(2, 132, 199)    # #0284C7 - Accent Brand Blue
C_CYAN_LIGHT   = RGBColor(240, 249, 255)  # #F0F9FF - Card Light Cyan Fill
C_SLATE_LIGHT  = RGBColor(248, 250, 252)  # #F8FAFC - Card Clean Slate Fill
C_AMBER_LIGHT  = RGBColor(255, 251, 235)  # #FFFBEB - Warning Card Fill
C_AMBER_BORDER = RGBColor(245, 158, 11)   # #F59E0B - Amber Accent
C_RED_ACCENT   = RGBColor(220, 38, 38)    # #DC2626 - Alert Red
C_GREEN_ACCENT = RGBColor(22, 163, 74)    # #16A34A - Safe Green
C_GREEN_LIGHT  = RGBColor(240, 253, 244)  # #F0FDF4 - Green Card Fill
C_BORDER_MUTED = RGBColor(203, 213, 225)  # #CBD5E1 - Crisp Outline
C_BORDER_BLUE  = RGBColor(186, 230, 253)  # #BAE6FD - Blue Outline
C_TEXT_DARK    = RGBColor(30, 41, 59)     # Slate 800
C_TEXT_MUTED   = RGBColor(71, 85, 105)    # Slate 600
C_WHITE        = RGBColor(255, 255, 255)

FONT_FAMILY = "Segoe UI"

def set_font(run, size_pt=11, bold=False, color=C_TEXT_DARK):
    run.font.name = FONT_FAMILY
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = color

def create_card(slide, left, top, width, height, fill_color=C_SLATE_LIGHT, border_color=C_BORDER_MUTED, border_width=1.0):
    """Creates a modern rounded rectangle container card."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = fill_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(border_width)
    else:
        card.line.fill.background()
    return card

def main():
    prs = Presentation(TEMPLATE_PATH)
    print(f"[*] Loaded template with {len(prs.slides)} slides.")

    # -------------------------------------------------------------
    # SLIDE 1: TITLE PAGE
    # -------------------------------------------------------------
    slide1 = prs.slides[0]
    
    # Clean and style existing title text
    for shape in list(slide1.shapes):
        if not shape.has_text_frame:
            continue
        text = shape.text_frame.text.strip()
        
        if "TITLE PAGE" in text:
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "SMART INDIA HACKATHON 2026"
            set_font(p.runs[0], size_pt=16, bold=True, color=C_SKY_BLUE)
            
        elif "SMART INDIA HACKATHON 2026" in text:
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "IDEA PRESENTATION"
            set_font(p.runs[0], size_pt=24, bold=True, color=C_NAVY_DARK)
            
        elif "Problem Statement ID" in text:
            # We will replace TextBox 9 with our custom designed executive card
            sp_elem = shape._element
            sp_elem.getparent().remove(sp_elem)

    # Add Project Hero Card on the left of Slide 1
    card1 = create_card(slide1, Inches(0.6), Inches(1.8), Inches(6.8), Inches(4.7), fill_color=C_SLATE_LIGHT, border_color=C_SKY_BLUE, border_width=1.5)
    tf1 = card1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = Inches(0.3)
    tf1.margin_right = Inches(0.3)
    tf1.margin_top = Inches(0.25)
    tf1.margin_bottom = Inches(0.2)

    # Project Banner inside Card
    p_head = tf1.paragraphs[0]
    p_head.text = "VIGIL-FLOOD"
    set_font(p_head.runs[0], size_pt=22, bold=True, color=C_NAVY_DARK)

    p_sub = tf1.add_paragraph()
    p_sub.text = "Multi-Source AI Flash Flood Early Warning & Tactical Decision Support System"
    set_font(p_sub.runs[0], size_pt=11, bold=True, color=C_SKY_BLUE)
    p_sub.space_after = Pt(12)

    metadata = [
        ("Problem Statement ID:", "SIH 26192"),
        ("Problem Statement Title:", "Flash Flood Prediction System for Hilly Regions using Multi-Source Data"),
        ("Theme:", "Disaster Management (Mountain Hydrology & Slope Stability)"),
        ("PS Category:", "Software (AI / 3D GIS Digital Twin / Edge IoT)"),
        ("Target Catchments:", "Beas & Sutlej River Basins (Samej Rampur, Mandi, Kullu - HP)"),
        ("Team Name / ID:", "[Your Registered Team Name / Team ID]")
    ]

    for label, val in metadata:
        p = tf1.add_paragraph()
        p.space_after = Pt(4)
        r_lbl = p.add_run()
        r_lbl.text = f"{label} "
        set_font(r_lbl, size_pt=10.5, bold=True, color=C_NAVY_DARK)
        r_val = p.add_run()
        r_val.text = val
        set_font(r_val, size_pt=10.5, bold=False, color=C_TEXT_DARK)

    # Add a pill badge at bottom of Slide 1 card
    badge = create_card(slide1, Inches(0.9), Inches(5.7), Inches(6.2), Inches(0.55), fill_color=C_CYAN_LIGHT, border_color=C_SKY_BLUE, border_width=1.0)
    tf_b = badge.text_frame
    tf_b.word_wrap = True
    p_b = tf_b.paragraphs[0]
    p_b.alignment = PP_ALIGN.CENTER
    r_b = p_b.add_run()
    r_b.text = "🛡️ 30-60 Min Actionable Lead Time | 96% AI Confidence | Zero-Cost Satellite Streams"
    set_font(r_b, size_pt=10, bold=True, color=C_SKY_BLUE)

    print("[+] Slide 1: Title Page populated.")

    # -------------------------------------------------------------
    # SLIDE 2: PROPOSED SOLUTION & CORE INNOVATIONS
    # -------------------------------------------------------------
    slide2 = prs.slides[1]
    
    # Configure Slide Title & remove placeholder text box
    for shape in list(slide2.shapes):
        if not shape.has_text_frame:
            continue
        t = shape.text_frame.text.strip()
        if "IDEA TITLE" in t:
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "PROPOSED SOLUTION: VIGIL-FLOOD PARADIGM"
            set_font(p.runs[0], size_pt=18, bold=True, color=C_NAVY_DARK)
        elif "Proposed Solution" in t or "Detailed explanation" in t:
            sp_elem = shape._element
            sp_elem.getparent().remove(sp_elem)

    # 4 Modern Feature Cards in a 2x2 Layout
    card_w = Inches(5.8)
    card_h = Inches(2.25)
    row1_t = Inches(1.3)
    row2_t = Inches(3.7)
    col1_l = Inches(0.6)
    col2_l = Inches(6.7)

    features = [
        (col1_l, row1_t, "🏘️ Hyper-Local Village & Ward Resolution (50m–100m)", C_CYAN_LIGHT, C_SKY_BLUE, [
            ("Overcomes District Forecast Limits: ", "Replaces broad district warnings with precise sub-catchment hazard zonation for vulnerable mountain habitations."),
            ("Dynamic Citizen Escape Windows: ", "Calculates 30–60 min actionable evacuation lead time tracking water rise rate (dH/dt) and distance to stream.")
        ]),
        (col2_l, row1_t, "🌊 Coupled Dual-Hazard AI Predictive Engine", C_SLATE_LIGHT, C_BORDER_MUTED, [
            ("Simultaneous Hazard Modeling: ", "Decouples river channel overtopping surge from steep hillside debris flows (Infinite Slope Factor of Safety Fs < 1.0)."),
            ("Production Ensemble: ", "Trained on 280,000+ real hourly hydrometeorological records across Himachal Pradesh; delivers 96% model confidence.")
        ]),
        (col1_l, row2_t, "🧠 Explainable AI (TreeSHAP) & What-If Sandbox", C_SLATE_LIGHT, C_BORDER_MUTED, [
            ("Transparent Decision Drivers: ", "Replaces black-box AI with game-theoretic attribution (+40% History, +34% Slope, +15% River Stage, +11% Soil)."),
            ("Real-Time Triage Sandbox: ", "Allows Incident Commanders to simulate cloudburst rainfall spikes (+40mm) and dam gate releases in <50ms.")
        ]),
        (col2_l, row2_t, "🚨 Multi-Modal Evacuation & NDMA CAP-SACHET", C_GREEN_LIGHT, C_GREEN_ACCENT, [
            ("ITU-T X.1303 Alerting: ", "Native XML generation compliant with NDMA CAP-SACHET and multi-lingual Cell Broadcast SMS to Panchayat Pradhans."),
            ("Uphill Shelter Allocation: ", "Directs villagers strictly uphill to pre-surveyed reinforced structures (e.g. Rampur College at +110m above floodline).")
        ])
    ]

    for left, top, header, fill, border, bullets in features:
        c = create_card(slide2, left, top, card_w, card_h, fill_color=fill, border_color=border, border_width=1.2)
        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_right = Inches(0.2)
        tf.margin_top = Inches(0.18)
        tf.margin_bottom = Inches(0.15)
        
        p0 = tf.paragraphs[0]
        p0.text = header
        set_font(p0.runs[0], size_pt=11.5, bold=True, color=C_NAVY_DARK)
        p0.space_after = Pt(6)

        for b_title, b_desc in bullets:
            p = tf.add_paragraph()
            p.space_after = Pt(3)
            r_t = p.add_run()
            r_t.text = "• " + b_title
            set_font(r_t, size_pt=10, bold=True, color=C_SKY_BLUE)
            r_d = p.add_run()
            r_d.text = b_desc
            set_font(r_d, size_pt=9.5, bold=False, color=C_TEXT_DARK)

    # Bottom KPI Ribbon
    ribbon = create_card(slide2, Inches(0.6), Inches(6.05), Inches(11.9), Inches(0.45), fill_color=C_NAVY_DARK, border_color=None)
    tf_r = ribbon.text_frame
    p_r = tf_r.paragraphs[0]
    p_r.alignment = PP_ALIGN.CENTER
    r_r = p_r.add_run()
    r_r.text = "⚡ KEY SYSTEM METRICS:   30–60 Min Actionable Lead Time   |   96% AI Ensemble Confidence   |   <50ms Inference Latency   |   Zero Recurring API Costs"
    set_font(r_r, size_pt=9.5, bold=True, color=C_WHITE)

    print("[+] Slide 2: Proposed Solution populated.")

    # -------------------------------------------------------------
    # SLIDE 3: TECHNICAL APPROACH & SYSTEM ARCHITECTURE
    # -------------------------------------------------------------
    slide3 = prs.slides[2]
    
    for shape in list(slide3.shapes):
        if not shape.has_text_frame:
            continue
        t = shape.text_frame.text.strip()
        if "TECHNICAL APPROACH" in t:
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "TECHNICAL APPROACH & SYSTEM ARCHITECTURE"
            set_font(p.runs[0], size_pt=18, bold=True, color=C_NAVY_DARK)
        elif "Technologies to be used" in t or "Methodology" in t:
            sp_elem = shape._element
            sp_elem.getparent().remove(sp_elem)

    # Embed High-Resolution Unified Slide Architecture Graphic
    if os.path.exists(UNIFIED_SLIDE_IMG):
        # Place image cleanly on Slide 3 (16:9 ratio, leaves room for bottom tech badges)
        slide3.shapes.add_picture(UNIFIED_SLIDE_IMG, Inches(0.6), Inches(1.2), Inches(8.3), Inches(4.67))
        print("[+] Embedded unified architecture graphic on Slide 3.")

    # Right-side Tech Stack Breakdown Card
    card_tech = create_card(slide3, Inches(9.1), Inches(1.2), Inches(3.4), Inches(4.67), fill_color=C_SLATE_LIGHT, border_color=C_SKY_BLUE, border_width=1.2)
    tf_tech = card_tech.text_frame
    tf_tech.word_wrap = True
    tf_tech.margin_left = Inches(0.18)
    tf_tech.margin_right = Inches(0.18)
    tf_tech.margin_top = Inches(0.18)

    p_th = tf_tech.paragraphs[0]
    p_th.text = "🛠️ TECHNOLOGY STACK"
    set_font(p_th.runs[0], size_pt=11.5, bold=True, color=C_NAVY_DARK)
    p_th.space_after = Pt(8)

    tech_specs = [
        ("📡 Ingestion & Telemetry:", "IMD Radar, NASA GPM (0.1°), GEE 30m DEM, IoT River Gauges, Soil Piezometers"),
        ("⚡ AI Engine & ML Core:", "XGBoost + LightGBM Ensemble, TreeSHAP Kernel, 280k real training samples"),
        ("📐 Physical Hydrology:", "D8 Flow Accumulation, Kinematic Wavefront Routing, Slope >35° Failure Zonation"),
        ("🖥️ Tactical GIS Digital Twin:", "MapLibre GL 3D Terrain, Leaflet 2D Triage, Three.js Valley Mesh, WebSockets (<500ms)"),
        ("📢 Alert Gateway:", "NDMA CAP-SACHET XML, Fast2SMS Telecom Multi-lingual Broadcast")
    ]

    for label, desc in tech_specs:
        p = tf_tech.add_paragraph()
        p.space_after = Pt(5)
        r_l = p.add_run()
        r_l.text = f"{label}\n"
        set_font(r_l, size_pt=9.5, bold=True, color=C_SKY_BLUE)
        r_d = p.add_run()
        r_d.text = desc
        set_font(r_d, size_pt=8.5, bold=False, color=C_TEXT_DARK)

    # Bottom workflow badge
    ribbon3 = create_card(slide3, Inches(0.6), Inches(6.0), Inches(11.9), Inches(0.45), fill_color=C_CYAN_LIGHT, border_color=C_SKY_BLUE)
    tf_r3 = ribbon3.text_frame
    p_r3 = tf_r3.paragraphs[0]
    p_r3.alignment = PP_ALIGN.CENTER
    r_r3 = p_r3.add_run()
    r_r3.text = "🔄 SEAMLESS DATA PIPELINE:   Satellite Telemetry  ➔  Hydrologic Physics  ➔  TreeSHAP ML Ensemble  ➔  CAP-SACHET Evacuation Alert"
    set_font(r_r3, size_pt=9.5, bold=True, color=C_SKY_BLUE)

    print("[+] Slide 3: Technical Approach populated.")

    # -------------------------------------------------------------
    # SLIDE 4: FEASIBILITY, VIABILITY & SENSOR FAULT TOLERANCE
    # -------------------------------------------------------------
    slide4 = prs.slides[3]
    
    for shape in list(slide4.shapes):
        if not shape.has_text_frame:
            continue
        t = shape.text_frame.text.strip()
        if "FEASIBILITY AND VIABILITY" in t:
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "FEASIBILITY, VIABILITY & SENSOR FAULT TOLERANCE"
            set_font(p.runs[0], size_pt=18, bold=True, color=C_NAVY_DARK)
        elif "Analysis of the feasibility" in t or "Potential challenges" in t:
            sp_elem = shape._element
            sp_elem.getparent().remove(sp_elem)

    # Top Row: 3 Feasibility Pillar Cards
    p_w = Inches(3.8)
    p_h = Inches(2.2)
    p_top = Inches(1.3)

    pillars = [
        (Inches(0.6), "💰 Economic Feasibility", C_SLATE_LIGHT, C_BORDER_MUTED, [
            ("Zero Recurring Data Costs: ", "Built 100% on open-access scientific streams (Copernicus 30m DEM, NASA SMAP, Open-Meteo)."),
            ("Cost-Effective Scale: ", "Zero commercial API subscriptions; minimal server compute overhead for state-wide deployment.")
        ]),
        (Inches(4.65), "⚡ Technical Feasibility", C_SLATE_LIGHT, C_BORDER_MUTED, [
            ("Edge-Ready Architecture: ", "Lightweight compiled ML models (<2MB) deployable on low-power solar edge gateways (Raspberry Pi/ESP32)."),
            ("High-Speed Inference: ", "<50ms model execution allows instant alert dispatch even in remote mountain valleys.")
        ]),
        (Inches(8.7), "🏛️ Operational Viability", C_SLATE_LIGHT, C_BORDER_MUTED, [
            ("NDMA Standard Compliant: ", "Native alignment with NDMA CAP-SACHET guidelines ensures immediate adoption by District EOCs."),
            ("Zero Learning Curve: ", "Intuitive color-coded triage (Red/Orange/Yellow) designed for rapid incident commander decisioning.")
        ])
    ]

    for left, title, fill, border, bullets in pillars:
        c = create_card(slide4, left, p_top, p_w, p_h, fill_color=fill, border_color=border, border_width=1.2)
        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.18)
        tf.margin_right = Inches(0.18)
        tf.margin_top = Inches(0.15)
        
        p0 = tf.paragraphs[0]
        p0.text = title
        set_font(p0.runs[0], size_pt=11.5, bold=True, color=C_NAVY_DARK)
        p0.space_after = Pt(6)

        for b_title, b_desc in bullets:
            p = tf.add_paragraph()
            p.space_after = Pt(3)
            r_t = p.add_run()
            r_t.text = "• " + b_title
            set_font(r_t, size_pt=9.5, bold=True, color=C_SKY_BLUE)
            r_d = p.add_run()
            r_d.text = b_desc
            set_font(r_d, size_pt=9, bold=False, color=C_TEXT_DARK)

    # Bottom Spotlight Card: The Sensor Outage USP (Alert Amber Styling)
    usp_card = create_card(slide4, Inches(0.6), Inches(3.7), Inches(11.9), Inches(2.7), fill_color=C_AMBER_LIGHT, border_color=C_AMBER_BORDER, border_width=1.5)
    tf_usp = usp_card.text_frame
    tf_usp.word_wrap = True
    tf_usp.margin_left = Inches(0.25)
    tf_usp.margin_right = Inches(0.25)
    tf_usp.margin_top = Inches(0.2)

    p_uh = tf_usp.paragraphs[0]
    p_uh.text = "🚨 THE KILLER USP — AUTONOMOUS SENSOR-OUTAGE FALLBACK (ZERO-DOWNTIME GUARANTEE)"
    set_font(p_uh.runs[0], size_pt=12.5, bold=True, color=RGBColor(180, 83, 9))
    p_uh.space_after = Pt(8)

    usp_points = [
        ("The Critical Vulnerability in Mountain Catchments: ", "During violent Himalayan cloudbursts, surging rivers carry heavy boulders that frequently destroy physical water level sensors. In traditional early warning systems, destroyed sensors cause fatal system failure and zero warnings."),
        ("VIGIL-FLOOD Autonomous Fallback Architecture: ", "Our system continuously monitors sensor heartbeat. The moment a river gauge disconnects, VIGIL-FLOOD autonomously transfers inference to an 11-Feature Satellite & Terrain Fallback Model within <100ms with zero manual intervention."),
        ("Physics-Reconstructed Warning Accuracy: ", "Threat levels are reconstructed dynamically using upstream satellite precipitation accumulation, soil piezometer saturation, and 30m DEM slope runoff velocity—preserving actionable evacuation alerts with 94%+ confidence even when ground sensors are wiped out!")
    ]

    for u_title, u_desc in usp_points:
        p = tf_usp.add_paragraph()
        p.space_after = Pt(4)
        r_t = p.add_run()
        r_t.text = "• " + u_title
        set_font(r_t, size_pt=10, bold=True, color=C_NAVY_DARK)
        r_d = p.add_run()
        r_d.text = u_desc
        set_font(r_d, size_pt=9.5, bold=False, color=C_TEXT_DARK)

    print("[+] Slide 4: Feasibility & Viability populated.")

    # -------------------------------------------------------------
    # SLIDE 5: IMPACT AND BENEFITS
    # -------------------------------------------------------------
    slide5 = prs.slides[4]
    
    for shape in list(slide5.shapes):
        if not shape.has_text_frame:
            continue
        t = shape.text_frame.text.strip()
        if "IMPACT AND BENEFITS" in t:
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "IMPACT, BENEFITS & STRATEGIC VALUE"
            set_font(p.runs[0], size_pt=18, bold=True, color=C_NAVY_DARK)
        elif "Potential impact" in t or "Benefits of the solution" in t:
            sp_elem = shape._element
            sp_elem.getparent().remove(sp_elem)

    # 4 Structured Impact Quadrants
    impact_cards = [
        (col1_l, row1_t, "👥 Citizen Life Safety & Evacuation (Social Impact)", C_GREEN_LIGHT, C_GREEN_ACCENT, [
            ("Zero Casualty Mission: ", "Transforms abstract weather data into a minute-by-minute 30–60 min escape countdown for mountain residents."),
            ("High-Ridge Shelter Routing: ", "Directs villagers safely uphill to reinforced structures (e.g. Rampur College at +110m above surge line), avoiding lethal riverside traps.")
        ]),
        (col2_l, row1_t, "🦺 NDRF & First Responder Tactical Edge (Operational)", C_CYAN_LIGHT, C_SKY_BLUE, [
            ("Live Tactical Dossier: ", "Empowers District EOCs with real-time flood hydrographs, terrain slope cross-sections, and road vulnerability status."),
            ("Proactive Pre-positioning: ", "Enables dispatch of rescue boats, ambulances, and NDRF battalions hours before mountain highways get severed.")
        ]),
        (col1_l, row2_t, "🛣️ Infrastructure & Asset Protection (Economic Impact)", C_SLATE_LIGHT, C_BORDER_MUTED, [
            ("Highway Safeguards: ", "Automated alerts trigger proactive closures of critical national highways (NH-3 / NH-5), preventing vehicle submergence."),
            ("Hydropower Resilience: ", "Provides early reservoir inflow warnings to dam operators (e.g. Nathpa Jhakri, Larji) for controlled spillway releases.")
        ]),
        (col2_l, row2_t, "🌿 Environmental & Pan-Himalayan Scale (Environmental)", C_SLATE_LIGHT, C_BORDER_MUTED, [
            ("Landslide Scar Monitoring: ", "Tracks post-monsoon soil saturation and progressive slope creep across active geological fault zones."),
            ("Pan-India Portability: ", "Architected for frictionless expansion across Uttarakhand, Jammu & Kashmir, Sikkim, and the Western Ghats.")
        ])
    ]

    for left, top, header, fill, border, bullets in impact_cards:
        c = create_card(slide5, left, top, card_w, card_h, fill_color=fill, border_color=border, border_width=1.2)
        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_right = Inches(0.2)
        tf.margin_top = Inches(0.18)
        tf.margin_bottom = Inches(0.15)
        
        p0 = tf.paragraphs[0]
        p0.text = header
        set_font(p0.runs[0], size_pt=11.5, bold=True, color=C_NAVY_DARK)
        p0.space_after = Pt(6)

        for b_title, b_desc in bullets:
            p = tf.add_paragraph()
            p.space_after = Pt(3)
            r_t = p.add_run()
            r_t.text = "• " + b_title
            set_font(r_t, size_pt=10, bold=True, color=C_SKY_BLUE)
            r_d = p.add_run()
            r_d.text = b_desc
            set_font(r_d, size_pt=9.5, bold=False, color=C_TEXT_DARK)

    # Bottom Beneficiary Ribbon
    ribbon5 = create_card(slide5, Inches(0.6), Inches(6.05), Inches(11.9), Inches(0.45), fill_color=C_NAVY_DARK, border_color=None)
    tf_r5 = ribbon5.text_frame
    p_r5 = tf_r5.paragraphs[0]
    p_r5.alignment = PP_ALIGN.CENTER
    r_r5 = p_r5.add_run()
    r_r5.text = "🎯 PRIMARY BENEFICIARIES:   Vulnerable Mountain Villages   |   NDRF & SDRF Tactical Squads   |   District EOC Commanders   |   State Disaster Management Authorities"
    set_font(r_r5, size_pt=9.5, bold=True, color=C_WHITE)

    print("[+] Slide 5: Impact & Benefits populated.")

    # -------------------------------------------------------------
    # SLIDE 6: RESEARCH AND REFERENCES
    # -------------------------------------------------------------
    slide6 = prs.slides[5]
    
    for shape in list(slide6.shapes):
        if not shape.has_text_frame:
            continue
        t = shape.text_frame.text.strip()
        if "RESEARCH" in t:
            tf = shape.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "RESEARCH RIGOR, SCIENTIFIC STANDARDS & REFERENCES"
            set_font(p.runs[0], size_pt=18, bold=True, color=C_NAVY_DARK)
        elif "Details / Links" in t:
            sp_elem = shape._element
            sp_elem.getparent().remove(sp_elem)

    # 3 Horizontal Structured Panels for Research & Standards
    ref_panels = [
        (Inches(1.3), Inches(1.5), "🏛️ National Disaster Management Protocols & Guidelines", C_CYAN_LIGHT, C_SKY_BLUE, [
            ("NDMA CAP-SACHET Guidelines (2022): ", "Standardized Common Alerting Protocol (ITU-T X.1303) compliant XML schema for multi-lingual citizen broadcast."),
            ("ISRO National Landslide Atlas (NRSC 2023): ", "Landslide hazard zonation methodology and historical scar spatial distribution across Himalayan river basins."),
            ("CWC & IMD Standard Operating Procedures: ", "Integration of Central Water Commission gauge thresholds and IMD Doppler Radar cloudburst detection criteria.")
        ]),
        (Inches(2.9), Inches(1.5), "🔬 Peer-Reviewed Hydrologic & Geotechnical Formulations", C_SLATE_LIGHT, C_BORDER_MUTED, [
            ("Kinematic Wave Wavefront Routing: ", "1D hydrodynamic flood surge routing (∂A/∂t + ∂Q/∂x = q) for steep, bedrock-confined mountain streams."),
            ("Infinite Slope Factor of Safety (Fs): ", "Limit equilibrium slope stability model accounting for transient pore water pressure rise during cloudburst storms."),
            ("TreeSHAP Explainable AI (Lundberg et al., Nature MI 2020): ", "Consistent game-theoretic Shapley value attribution for tree ensemble hazard predictions.")
        ]),
        (Inches(4.5), Inches(1.5), "🛰️ Open-Access Scientific Geospatial Data Streams", C_SLATE_LIGHT, C_BORDER_MUTED, [
            ("Copernicus GLO-30 DEM: ", "European Space Agency 30m high-resolution digital elevation model for watershed flow accumulation (D8)."),
            ("NASA GPM IMERG & SMAP: ", "Global Precipitation Measurement 0.1° satellite rainfall and Soil Moisture Active Passive root-zone moisture."),
            ("ECMWF ERA5-Land & Open-Meteo: ", "Hourly atmospheric reanalysis parameters powering low-latency sub-catchment meteorological ingestion.")
        ])
    ]

    for top, height, title, fill, border, bullets in ref_panels:
        c = create_card(slide6, Inches(0.6), top, Inches(11.9), height, fill_color=fill, border_color=border, border_width=1.2)
        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_right = Inches(0.2)
        tf.margin_top = Inches(0.12)
        tf.margin_bottom = Inches(0.1)

        p0 = tf.paragraphs[0]
        p0.text = title
        set_font(p0.runs[0], size_pt=11, bold=True, color=C_NAVY_DARK)
        p0.space_after = Pt(4)

        for b_title, b_desc in bullets:
            p = tf.add_paragraph()
            p.space_after = Pt(2)
            r_t = p.add_run()
            r_t.text = "• " + b_title
            set_font(r_t, size_pt=9.5, bold=True, color=C_SKY_BLUE)
            r_d = p.add_run()
            r_d.text = b_desc
            set_font(r_d, size_pt=9, bold=False, color=C_TEXT_DARK)

    # Bottom Citation Banner
    ribbon6 = create_card(slide6, Inches(0.6), Inches(6.1), Inches(11.9), Inches(0.4), fill_color=C_NAVY_DARK, border_color=None)
    tf_r6 = ribbon6.text_frame
    p_r6 = tf_r6.paragraphs[0]
    p_r6.alignment = PP_ALIGN.CENTER
    r_r6 = p_r6.add_run()
    r_r6.text = "📚 All hydrological models and AI architectures are fully documented, open-source compliant, and verified against real Himalayan disaster events."
    set_font(r_r6, size_pt=9, bold=True, color=C_WHITE)

    print("[+] Slide 6: Research & References populated.")

    # -------------------------------------------------------------
    # SLIDE 7: DELETE TEMPLATE INSTRUCTION SLIDE
    # -------------------------------------------------------------
    if len(prs.slides) >= 7:
        rId = prs.slides._sldIdLst[6].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[6]
        print("[+] Slide 7 (Instruction Slide) successfully removed per SIH guidelines.")

    # Save output
    prs.save(OUTPUT_PATH)
    print(f"\n[SUCCESS] Final SIH Presentation saved to: {OUTPUT_PATH}")
    print(f"Total Slides: {len(prs.slides)}")

if __name__ == "__main__":
    main()
